"""Text extraction for the indexable file types.

Every extractor returns plain text or raises ExtractionError. Extraction is
deliberately best-effort and never fatal to an index run: one corrupt PDF in a
folder must not stop the other 500 files from being indexed.
"""
from pathlib import Path

SUPPORTED_SUFFIXES = {".txt", ".md", ".pdf", ".docx", ".xlsx"}


class ExtractionError(Exception):
    pass


def _extract_txt(path: Path) -> str:
    # errors="replace" so one bad byte doesn't lose the whole file.
    return path.read_text(encoding="utf-8", errors="replace")


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")  # many PDFs are "encrypted" with an empty password
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError(f"encrypted PDF: {exc}") from exc
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    # Tables carry a lot of the real content in practice; paragraphs alone miss them.
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    # read_only + data_only: stream large sheets, and read cached formula RESULTS
    # rather than the formula text ("=SUM(B2:B9)" is not searchable content).
    wb = load_workbook(str(path), read_only=True, data_only=True)
    lines = []
    try:
        for sheet in wb.worksheets:
            lines.append(f"# sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append("\t".join(cells))
    finally:
        wb.close()
    return "\n".join(lines)


_EXTRACTORS = {
    ".txt": _extract_txt,
    ".md": _extract_txt,
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
}


def extract_text(path: Path) -> str:
    """Extract plain text from a supported file. Raises ExtractionError."""
    path = Path(path)
    fn = _EXTRACTORS.get(path.suffix.lower())
    if fn is None:
        raise ExtractionError(f"unsupported file type: {path.suffix!r}")
    try:
        return fn(path).strip()
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"{type(exc).__name__}: {exc}") from exc


def chunk_text(text: str, chunk_chars: int = 1200, overlap: int = 200) -> list[str]:
    """Split into overlapping chunks, preferring to break at paragraph/line ends.

    Overlap matters for retrieval: a sentence straddling a boundary would
    otherwise be split across two chunks and match neither well.
    """
    text = text.strip()
    if not text:
        return []
    if len(text) <= chunk_chars:
        return [text]

    chunks, start = [], 0
    while start < len(text):
        end = min(start + chunk_chars, len(text))
        if end < len(text):
            # Back off to the nearest paragraph or line break in the last 30%.
            window = text[start:end]
            for sep in ("\n\n", "\n", ". "):
                idx = window.rfind(sep, int(chunk_chars * 0.7))
                if idx != -1:
                    end = start + idx + len(sep)
                    break
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [c for c in chunks if c]
