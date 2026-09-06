"""Create realistic sample documents in the sandbox for indexer testing."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
from config import ALLOWED_ROOTS  # noqa: E402

ROOT = ALLOWED_ROOTS[0]


def make_txt():
    (ROOT / "notes").mkdir(parents=True, exist_ok=True)
    (ROOT / "notes" / "meeting_notes.txt").write_text(
        "Team sync - 14 March 2026\n\n"
        "Attendees: Priya, Marcus, Dana\n\n"
        "- Marcus reported the Helsinki warehouse migration slipped two weeks. "
        "Root cause was a delayed customs clearance on the racking shipment.\n"
        "- Priya raised that the quarterly budget spreadsheet still lists the old "
        "freight vendor. Needs updating before the board review.\n"
        "- Dana will draft the renewal terms for the Contoso contract. "
        "Target date is 2 April.\n\n"
        "Action items:\n"
        "1. Update freight vendor in Q1 budget (Priya)\n"
        "2. Reforecast Helsinki go-live (Marcus)\n"
        "3. Draft Contoso renewal (Dana)\n",
        encoding="utf-8",
    )
    (ROOT / "notes" / "ideas.md").write_text(
        "# Product ideas\n\n"
        "## Offline-first sync\n"
        "Customers in low-connectivity regions keep losing edits. An offline "
        "queue with conflict resolution would remove the top support complaint.\n\n"
        "## Warehouse capacity dashboard\n"
        "Ops currently tracks pallet capacity in a spreadsheet emailed weekly. "
        "A live dashboard reading from the WMS would replace it.\n",
        encoding="utf-8",
    )


def make_pdf():
    from reportlab.lib.pagesizes import LETTER
    from reportlab.pdfgen import canvas

    (ROOT / "reports").mkdir(parents=True, exist_ok=True)
    path = ROOT / "reports" / "annual_review.pdf"
    c = canvas.Canvas(str(path), pagesize=LETTER)
    lines_p1 = [
        "Annual Operations Review 2025",
        "",
        "Executive summary",
        "Total shipment volume grew 18% year over year, driven mainly by the",
        "Nordic corridor. Cost per pallet fell from 42.10 EUR to 38.75 EUR",
        "after the freight vendor consolidation completed in Q3.",
        "",
        "Risks",
        "The Helsinki warehouse migration remains the largest open risk. A two",
        "week slip has already been absorbed; further delay would push go-live",
        "past the peak season freeze in November.",
    ]
    lines_p2 = [
        "Appendix A - Vendor summary",
        "",
        "Northwind Freight was retained as primary carrier through 2027.",
        "Contoso Logistics remains the secondary carrier; that contract is up",
        "for renewal on 2 April 2026 and has not yet been drafted.",
    ]
    for lines in (lines_p1, lines_p2):
        y = 720
        for line in lines:
            c.setFont("Helvetica-Bold" if y == 720 else "Helvetica", 14 if y == 720 else 11)
            c.drawString(72, y, line)
            y -= 18
        c.showPage()
    c.save()


def make_docx():
    import docx

    doc = docx.Document()
    doc.add_heading("Contoso Logistics - Renewal Terms (DRAFT)", level=1)
    doc.add_paragraph(
        "This draft sets out the proposed renewal terms for the secondary "
        "carrier agreement with Contoso Logistics, expiring 2 April 2026."
    )
    doc.add_heading("Commercial terms", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Current"
    table.rows[0].cells[2].text = "Proposed"
    for item, cur, prop in [
        ("Base rate per pallet", "38.75 EUR", "36.90 EUR"),
        ("Fuel surcharge cap", "6.0%", "4.5%"),
        ("Volume commitment", "12,000 pallets", "15,000 pallets"),
        ("Notice period", "90 days", "60 days"),
    ]:
        row = table.add_row()
        row.cells[0].text, row.cells[1].text, row.cells[2].text = item, cur, prop
    doc.add_paragraph(
        "Dana Whitfield is the responsible owner for this draft. Legal review "
        "has not yet been scheduled."
    )
    doc.save(str(ROOT / "reports" / "contoso_renewal_draft.docx"))


def make_xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Q1 Budget"
    ws.append(["Category", "Budgeted", "Actual", "Variance"])
    for row in [
        ["Freight - Northwind", 412000, 398500, "=C2-B2"],
        ["Freight - Contoso", 118000, 131200, "=C3-B3"],
        ["Warehouse lease", 96000, 96000, "=C4-B4"],
        ["Helsinki migration", 240000, 287400, "=C5-B5"],
        ["Contract labour", 74500, 69100, "=C6-B6"],
    ]:
        ws.append(row)
    ws2 = wb.create_sheet("Notes")
    ws2.append(["Freight vendor listed for Contoso is out of date - update before board review"])
    ws2.append(["Helsinki migration is over budget by 47,400 due to customs delay"])
    wb.save(str(ROOT / "reports" / "q1_budget.xlsx"))


if __name__ == "__main__":
    make_txt()
    make_pdf()
    make_docx()
    make_xlsx()
    for p in sorted(ROOT.rglob("*")):
        if p.is_file():
            print(f"  {p.stat().st_size:8,d}  {p.relative_to(ROOT)}")
    print(f"\nSandbox ready at {ROOT}")
